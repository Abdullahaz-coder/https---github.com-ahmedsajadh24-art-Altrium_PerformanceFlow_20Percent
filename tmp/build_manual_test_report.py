from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "documents" / "Altrium_PerformanceFlow_Completed_Manual_Test_Cases_2026-09-23.docx"

# The template supplied by the user contains these six columns. The detailed
# records below keep that schema while adding evidence-based execution notes.
CASES = [
    ("TC01", "Employee profile", "As HR, create a Finance employee and choose a supervisor from Operations; then retry with an Operations employee.", "Cross-department reporting is rejected; matching department is accepted.", "Mismatched assignment was not saved. Matching Operations profile was saved.", "Pass"),
    ("TC02", "Review assignment", "Create an employee without an active blueprint. Try to assign them to a cycle, then add an active blueprint item and retry.", "Only the employee with an active blueprint can be assigned.", "Assignment was blocked without an active item and saved after activation.", "Pass"),
    ("TC03", "HR records", "Search the centralized records as HR, then open the same route as another role.", "HR can retrieve records; unauthorized roles cannot.", "HR search returned records. Other roles were denied access.", "Pass"),
    ("TC04", "Reminders", "Trigger reminders twice for the same outstanding action.", "One notification is delivered per action, without duplicates.", "The outstanding action produced one reminder only.", "Pass"),
    ("TC05", "PDP workflow", "Record a PAR outcome requiring a PDP; create the plan, update progress, and open monitoring.", "Plan, employee progress, and supervisor/HR monitoring remain connected.", "PDP creation, progress update, and monitoring states advanced correctly.", "Pass"),
    ("TC06", "Cycle activation", "Activate a configured cycle containing an assigned employee.", "Each workflow stage receives one appropriate action.", "One action per stage was created; no duplicate stage actions appeared.", "Pass"),
    ("TC07", "Self assessment", "Send malformed self-assessment data to the save endpoint.", "Invalid request is rejected without a server error.", "The endpoint returned HTTP 400 and did not crash.", "Pass"),
    ("TC08", "Supervisor input", "Submit an invalid supervisor value while editing an employee.", "Input is rejected safely; the app remains available.", "Invalid supervisor input was handled without a server error.", "Pass"),
    ("TC09", "Staff creation", "As HR, create a Supervisor account using the profile form.", "Supervisor login and profile are created with the selected role.", "Supervisor account and matching profile were created.", "Pass"),
    ("TC10", "Cycle closure", "Attempt closure before PAR outcome and acknowledgement; complete both and retry.", "Closure is blocked until required final records exist.", "Premature closure was rejected; eligible closure succeeded.", "Pass"),
    ("TC11", "Manager reassignment", "As HR, reassign an outstanding approval to another manager.", "Pending decision moves to the newly assigned manager.", "Approval ownership and action were reassigned correctly.", "Pass"),
    ("TC12", "Change requests", "As Manager, request changes from two selected participants; inspect each recipient's view.", "Only selected people receive their own request and note.", "Two targeted requests were created and kept private to their recipients.", "Pass"),
    ("TC13", "Employee status", "As HR, set an employee inactive and inspect previous reviews.", "Employee becomes inactive but historical records remain.", "Inactive status saved; review history was retained.", "Pass"),
    ("TC14", "Status display", "Open the inactive employee workspace after a status update.", "Inactive status has warning styling and one success message.", "Warning status rendered and only one flash message appeared.", "Pass"),
    ("TC15", "Full review flow", "Create cycle and blueprint; run self, peer, supervisor and manager stages; acknowledge, hold PAR, create PDP, close.", "All handoffs complete in order and records remain available.", "The isolated end-to-end workflow reached a closed cycle successfully.", "Pass"),
    ("TC16", "Availability", "Submit availability values with mixed timezone formats.", "Invalid or mixed input is handled without an internal error.", "The availability endpoint completed without a crash.", "Pass"),
    ("TC17", "Role change", "Change an account role while its session remains active; request a protected workspace.", "The current role governs access immediately.", "The existing session reflected the new role permissions.", "Pass"),
    ("TC18", "PAR retention", "Close a review, then attempt to cancel its retained PAR meeting.", "Closed review meeting record cannot be cancelled.", "Cancellation was rejected and the historical meeting remained.", "Pass"),
    ("TC19", "PDP after closure", "Close the cycle and open the employee's PDP workspace.", "Development remains accessible after review closure.", "The development workspace remained reachable after closure.", "Pass"),
    ("TC20", "PDP privacy", "As one employee, request another employee's plan URL.", "The other person's plan is not disclosed.", "Access to the other employee's plan was denied.", "Pass"),
    ("TC21", "PAR scheduling", "Try a past meeting time and try marking an unheld meeting complete.", "Past scheduling and premature completion are rejected.", "Both invalid transitions were blocked.", "Pass"),
    ("TC22", "PAR conflict check", "Block an attendee's time, then try the conflicting meeting slot; inspect notifications.", "Conflicting slot is rejected and working link is private to attendees.", "Conflict was rechecked on save; attendee notification used a private link.", "Pass"),
    ("TC23", "Navigation", "Open role navigation as HR, Supervisor, Manager and Employee.", "Allowed destinations open and only one item is active.", "All tested role routes were reachable with unique active selection.", "Pass"),
    ("TC24", "Password change", "Change an account password, then reuse an existing session.", "Password changes and old sessions are revoked.", "The previous session stopped working after the password change.", "Pass"),
    ("TC25", "PDP edit", "Edit a plan after employee progress and reminder references exist.", "Plan edits retain progress and linked reminders.", "Progress records and reminder references remained intact.", "Pass"),
    ("TC26", "PDP dates", "Submit invalid plan and activity dates.", "Invalid dates are rejected without a partial write.", "Invalid date input did not update stored plan data.", "Pass"),
    ("TC27", "PDP display", "Open a plan as HR.", "Heading identifies an employee plan, not HR's own plan.", "HR view did not display the misleading 'My plan' heading.", "Pass"),
    ("TC28", "Journey status", "Advance the cohort through review stages and inspect HR dashboard.", "Journey stage reflects actual active-cohort progress.", "The displayed stage followed the active cohort state.", "Pass"),
    ("TC29", "Account inactivity", "Deactivate an account; attempt sign-in and reuse an existing session.", "Inactive account cannot sign in or continue a session.", "New sign-in and existing session access were blocked.", "Pass"),
    ("TC30", "CSRF protection", "Submit a protected form without token, then with a valid token.", "Missing token is rejected and valid token accepted.", "Missing token was rejected; valid-token request succeeded.", "Pass"),
    ("TC31", "PDP transaction", "Force an action-sync failure while saving a PDP edit.", "Plan change rolls back instead of leaving inconsistent data.", "The attempted edit was rolled back with no partial update.", "Pass"),
    ("TC32", "Database integrity", "Open database connections and inspect foreign-key enforcement.", "Foreign keys are enabled for application connections.", "Application connections enforced SQLite foreign keys.", "Pass"),
    ("TC33", "Cross-browser UI", "Repeat the complete workflow in Chrome, Edge and Firefox at desktop width.", "Controls, layouts and validation work consistently in all browsers.", "Not run in separate browser environments during this execution.", "Not run"),
    ("TC34", "Responsive UI", "Open every major screen at phone, tablet and desktop widths; inspect overflow and dialogs.", "Text remains readable; no clipping, overlap or inaccessible action.", "Not run as a full device-size visual sweep during this execution.", "Not run"),
    ("TC35", "Accessibility", "Navigate each key workflow by keyboard and inspect labels, focus order and contrast.", "All primary actions are keyboard-accessible and clearly labeled.", "Not run as a full assisted-technology accessibility audit.", "Not run"),
    ("TC36", "Performance", "Measure dashboard, cycle, PAR and PDP response times under representative concurrent users.", "Response time stays within the agreed service target.", "Not run; no agreed production load target or staging load test was available.", "Not run"),
    ("TC37", "Recovery", "Restore a production-like backup and verify review, PAR and PDP data consistency.", "Restored records are complete and access controls remain intact.", "Not run; production-like backup/restore was outside this isolated test run.", "Not run"),
]

SECTIONS = [
    ("Functional security and workflow cases", 0, 32),
    ("Non functional checks awaiting execution", 32, 37),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        bd.append(el)
    tc_pr.append(bd)


def margin(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, value in (("top", "75"), ("bottom", "75"), ("left", "75"), ("right", "75")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), value)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.6)
sec.left_margin = Inches(0.43)
sec.right_margin = Inches(0.43)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(4)
for name, size in (("Title", 18), ("Heading 1", 12), ("Heading 2", 10)):
    style = styles[name]
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(5)

# Word's built-in Title style may carry a blue bottom rule; the submission
# design relies on whitespace only.
title_ppr = styles["Title"]._element.get_or_add_pPr()
for child in list(title_ppr):
    if child.tag == qn("w:pBdr"):
        title_ppr.remove(child)

doc.add_paragraph("Altrium PerformanceFlow Manual Test Cases", style="Title")
doc.add_paragraph("Sprint 1 and Sprint 2 functional and non functional verification")
p = doc.add_paragraph()
p.add_run("Prepared date: ").bold = True
p.add_run("23 September 2026    ")
p.add_run("Result: ").bold = True
p.add_run("32 Pass | 0 Fail | 5 Not run")

doc.add_heading("Purpose and execution basis", 1)
doc.add_paragraph(
    "This record follows the supplied test case template: test number, feature, steps, expected outcome, "
    "actual outcome and status. It covers the implemented review lifecycle from employee setup through PAR meetings, "
    "development planning and cycle closure. The 32 passed outcomes were verified on 23 September 2026 by running "
    "the application's isolated regression suite against a temporary database copy. They are route-level and data-level "
    "checks, not a claim that each step was repeated by a human in a browser. Five environment-dependent checks remain "
    "explicitly marked Not run. The live application database was not changed by this test run."
)

doc.add_heading("Test environment and data", 1)
doc.add_paragraph(
    "Environment: local Flask application and SQLite test copy. Roles: HR, Supervisor, Manager and Employee. "
    "Representative records: 2026 review cycle, active employee profiles, performance blueprint items, self and peer "
    "reviews, manager approval, PAR meeting, PDP and notifications. Regression command: unittest discovery across "
    "test_regressions.py and test_role_audit.py. Execution result: 32 tests ran successfully in about six seconds."
)
doc.add_paragraph(
    "Status meaning: Pass = the expected result was observed in the isolated test run; Fail = observed result differs; "
    "Not run = no trustworthy actual result is available. Overall disposition: conditional pass for the verified "
    "functional and security scope; browser, accessibility, load and recovery acceptance still requires separate execution."
)

widths = [0.48, 0.88, 2.12, 1.57, 1.79, 0.73]
headers = ["Feature #", "Feature name", "Steps", "Expected outcome", "Actual outcome", "Status"]

for section_name, start, end in SECTIONS:
    doc.add_heading(section_name, 1)
    table = doc.add_table(rows=1, cols=6)
    table.autofit = False
    for i, (head, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = head
        shade(cell, "1F2B45")
        margin(cell)
        borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)

    for row_index, record in enumerate(CASES[start:end]):
        row = table.add_row()
        row_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        row_pr.append(no_split)
        for i, (value, width) in enumerate(zip(record, widths)):
            cell = row.cells[i]
            cell.width = Inches(width)
            cell.text = value
            margin(cell)
            borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                shade(cell, "F4F7FB")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.08
            if i in (0, 5):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8)
                if i in (0, 5):
                    run.font.bold = True
                if i == 5 and value == "Not run":
                    run.font.color.rgb = RGBColor(151, 94, 18)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Altrium PerformanceFlow | Test execution record | 23 September 2026")
for run in footer.runs:
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(95, 105, 122)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
